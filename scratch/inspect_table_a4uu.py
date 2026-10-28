import openpyxl

wb = openpyxl.load_workbook(r"D:\PainelNovo-main\Monitoramento\Monitoramento individual_A4UU_v4.xlsm", data_only=True)
ws = wb["Monitoramento"]

for tbl in ws.tables.values():
    print("Table name:", tbl.name, "range:", tbl.ref)
    col_names = [col.name for col in tbl.tableColumns]
    print("Table columns:", col_names)

# Let's inspect rows around row 25 with column names
min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(tbl.ref)
print(f"Header row is {min_row}")
table_headers = [ws.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
print("Actual headers in table range:", table_headers)

for r in range(min_row + 1, min_row + 100):
    vals = [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
    if any("CSC0823719" == str(v) for v in vals):
        print(f"\n--- Found row in Table at Excel row {r} ---")
        for h, v in zip(table_headers, vals):
            print(f"  {h}: {v}")
