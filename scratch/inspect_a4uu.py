import openpyxl

wb = openpyxl.load_workbook(r"D:\PainelNovo-main\Monitoramento\Monitoramento individual_A4UU_v4.xlsm", data_only=True)
print("Sheets in A4UU:", wb.sheetnames)
ws = wb["Monitoramento"] if "Monitoramento" in wb.sheetnames else wb.active

headers = [cell.value for cell in ws[1]]
print("Headers:", headers)

for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    row_str = " ".join([str(v) for v in row if v is not None])
    if "CSC0823719" in row_str or "000023/2026" in row_str:
        print(f"\n--- Found at row {row_idx} ---")
        for h, val in zip(headers, row):
            if val is not None or h in ["Contrato", "Situação", "Localidade", "UF", "Código OS", "Código da\nsolicitação", "Descrição da atividade"]:
                print(f"  {h!r}: {val!r}")
