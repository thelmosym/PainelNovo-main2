import openpyxl

wb = openpyxl.load_workbook(r"D:\PainelNovo-main\Painel de controle MemoriaPetrobras V82.xlsm", data_only=True)
print("Sheet names:", wb.sheetnames)

ws_dados = wb["DADOS"]
headers = [c.value for c in ws_dados[1]]
print("DADOS Headers:", headers)

# Find CSC0823719 or check row 661
print("\n--- Inspecting Row 661 of DADOS ---")
row_661 = [c.value for c in ws_dados[661]]
for h, v in zip(headers, row_661):
    print(f"  {h}: {v}")

# Search for CSC0823719 in all rows of DADOS
for r in range(2, ws_dados.max_row + 1):
    vals = [c.value for c in ws_dados[r]]
    if any("CSC0823719" == str(v) for v in vals):
        print(f"\n--- Found CSC0823719 at row {r} ---")
        for h, v in zip(headers, vals):
            print(f"  {h}: {v}")
