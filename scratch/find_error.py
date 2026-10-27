import openpyxl, subprocess

def test_chunk(build_fn, name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    build_fn(ws)
    fname = f"scratch/test_{name}.xlsx"
    wb.save(fname)
    cmd = [
        'powershell', '-Command',
        f'& {{ $xl = New-Object -ComObject Excel.Application; try {{ $wb = $xl.Workbooks.Open([System.IO.Path]::GetFullPath("{fname}")); Write-Host "{name} OK"; $wb.Close($false) }} catch {{ Write-Host "{name} ERROR: $($_.Exception.Message)" }} finally {{ $xl.Quit(); [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null }} }}'
    ]
    res = subprocess.run(cmd, capture_output=True, text=True, encoding='cp1252', errors='replace')
    print(res.stdout.strip())

# Test A: Title + KPI cards
def build_a(ws):
    ws.merge_cells("A1:N1")
    ws["A1"] = "Title"
    ws.merge_cells("A4:B5")
    ws["A4"] = "Card 1"

test_chunk(build_a, "part_a")

# Test B: Table 1 headers and data without formulas
def build_b(ws):
    ws.cell(row=8, column=1, value="Contrato")
    ws.cell(row=9, column=1, value="IRON-LT1")
    ws.cell(row=9, column=8, value=231)
    ws.cell(row=9, column=9, value=75)

test_chunk(build_b, "part_b")

# Test C: Table 1 formulas
def build_c(ws):
    ws.cell(row=9, column=8, value="=I9+J9")
    ws.cell(row=9, column=9, value=75)
    ws.cell(row=9, column=10, value=156)
    ws.cell(row=9, column=11, value="=J9/H9")
    ws.cell(row=9, column=12, value="=I9/H9*100")
    ws.cell(row=9, column=14, value="=IF(L9<90,0.97,IF(L9<95,0.98,IF(L9<99,0.99,1)))")

test_chunk(build_c, "part_c")

# Test D: Section 2 & 3
def build_d(ws):
    ws.cell(row=17, column=1, value="Faixa")
    ws.cell(row=18, column=1, value=">=99%")
    ws.cell(row=18, column=2, value=1.0)
    ws.cell(row=26, column=2, value="= 0,97") # NOTICE "= 0,97" starts with "=" !!!

test_chunk(build_d, "part_d")
