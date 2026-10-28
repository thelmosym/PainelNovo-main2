import openpyxl, datetime
from collections import defaultdict

# 1. Load holidays from PainelAgosto TabelaB-V2.xlsx
wb_tb = openpyxl.load_workbook(r'd:\PainelNovo-main\PainelAgosto\Tabelas Especificas\TabelaB-V2.xlsx', data_only=True)
ws_cal = wb_tb['Calendario']
feriados = []
for r in ws_cal.iter_rows(min_row=2, values_only=True):
    dt, desc, tipo, mun, uf = r[0], r[1], r[2], r[3], r[4]
    if dt:
        d = dt.date() if isinstance(dt, datetime.datetime) else dt
        feriados.append((d, (mun or '').strip().upper(), (uf or '').strip().upper()))

def tem_expediente(d, mun, uf):
    if d.weekday() >= 5: # Sat, Sun
        return False
    mun_str = (mun or '').strip().upper()
    uf_str = (uf or '').strip().upper()
    for f_d, f_m, f_u in feriados:
        if f_d == d:
            u_ok = (f_u == 'TODOS' or not f_u or f_u == uf_str)
            m_ok = (f_m == 'TODOS' or not f_m or f_m == mun_str)
            if u_ok and m_ok:
                return False
    return True

def add_business_hours(d_start, t_start, duration_hours, mun, uf):
    remaining = datetime.timedelta(hours=duration_hours)
    cur_d = d_start
    cur_t = t_start
    h_entry = datetime.time(8, 0)
    h_exit = datetime.time(17, 0)
    h_lunch_start = datetime.time(12, 0)
    h_lunch_end = datetime.time(13, 0)
    
    while remaining > datetime.timedelta(0):
        if not tem_expediente(cur_d, mun, uf):
            cur_d += datetime.timedelta(days=1)
            cur_t = datetime.time(0, 0)
            continue
        if cur_t > h_exit: cur_t = h_exit
        elif cur_t < h_entry: cur_t = h_entry
        elif h_lunch_start < cur_t < h_lunch_end: cur_t = h_lunch_end
        
        if cur_t < h_lunch_start:
            avail = datetime.datetime.combine(cur_d, h_lunch_start) - datetime.datetime.combine(cur_d, cur_t)
            if remaining <= avail:
                return datetime.datetime.combine(cur_d, cur_t) + remaining
            remaining -= avail
            cur_t = h_lunch_end
            
        avail_aft = datetime.datetime.combine(cur_d, h_exit) - datetime.datetime.combine(cur_d, cur_t)
        if remaining <= avail_aft:
            return datetime.datetime.combine(cur_d, cur_t) + remaining
        remaining -= avail_aft
        cur_d += datetime.timedelta(days=1)
        cur_t = datetime.time(0, 0)
    return datetime.datetime.combine(cur_d, cur_t)

def calc_sla_dt(prazo, d_abertura, mun, uf):
    d_ini = d_abertura.date()
    t_ini = d_abertura.time()
    if prazo == 0: return None
    elif prazo == 0.5:
        base = add_business_hours(d_ini, t_ini, 4, mun, uf)
        dbase = base.date()
        if datetime.time(8,0) <= t_ini <= datetime.time(11,0):
            return datetime.datetime.combine(dbase, datetime.time(17,0))
        elif datetime.time(11,0) < t_ini <= datetime.time(16,0):
            return add_business_hours(dbase, datetime.time(12,0), 8, mun, uf)
        else:
            return datetime.datetime.combine(dbase, datetime.time(17,0))
    elif prazo == 1: return add_business_hours(d_ini, t_ini, 8, mun, uf)
    elif prazo == 2: return add_business_hours(d_ini, t_ini, 16, mun, uf)
    elif prazo == 3: return add_business_hours(d_ini, t_ini, 24, mun, uf)
    elif prazo == 4: return add_business_hours(d_ini, t_ini, 32, mun, uf)
    elif prazo == 5: return add_business_hours(d_ini, t_ini, 40, mun, uf)
    return None

# Load Painel T2M from Painel Agosto.xlsm
wb_ag = openpyxl.load_workbook(r'd:\PainelNovo-main\PainelAgosto\Painel Agosto.xlsm', data_only=True)
ws = wb_ag['Painel T2M']
headers = [c.value for c in ws[1]]

def get_idx(name):
    for i, h in enumerate(headers):
        if h == name or (h and h.replace('\n', ' ') == name): return i
    return -1

c_contrato = get_idx('Contrato')
c_situacao = get_idx('Situação')
c_ref_fdm = get_idx('Ref. FDM')
c_ref_prazo = get_idx('Ref. prazo')
c_calc_fdm = get_idx('CalculoFDM')
c_desc = get_idx('Descrição da atividade')
c_d_ab = get_idx('D. abertura')
c_d_fc = get_idx('D. fechamento')
c_pz_comb = get_idx('Prazo combinado')
c_mun = get_idx('Município')
c_uf = get_idx('UF')
c_grupo = get_idx('Grupo.1')
c_prazo = get_idx('tabelaPrazoContrato.Prazo')
c_aplic = get_idx('Aplicação')

counts = defaultdict(lambda: {'TSPE1': 0, 'TSPR1': 0, 'TSPE2': 0, 'TSPR2': 0, 'total': 0})
status_dist = defaultdict(lambda: defaultdict(int))

for row_idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    contrato = r[c_contrato]
    situacao = r[c_situacao]
    ref_fdm = r[c_ref_fdm]
    ref_prazo = r[c_ref_prazo]
    calc_fdm = r[c_calc_fdm]
    desc = r[c_desc] or ''
    d_ab = r[c_d_ab]
    d_fc = r[c_d_fc]
    pz_comb = r[c_pz_comb]
    mun = r[c_mun]
    uf = r[c_uf]
    grupo = r[c_grupo]
    prazo = r[c_prazo]
    aplic = r[c_aplic]
    
    classif = None
    status_prazo = None
    log = ''
    
    if calc_fdm == 'S':
        classif = 'FDM definido'
    elif ref_fdm is None or ref_fdm == 0:
        classif = 'Isento'
    elif ref_prazo == 'Tabela':
        if not contrato: log = 'Contrato vazio; '
        elif not d_ab: log = 'D. abertura vazio; '
        elif not d_fc: log = 'D. fechamento vazio; '
        elif not mun: log = 'Municipio vazio; '
        elif not uf: log = 'UF vazio; '
        elif not grupo or grupo == 'Nao encontrado': log = 'Municipio sem grupo valido; '
        elif not prazo or prazo == 0: log = 'FDM nao permitido; '
        else:
            sla_dt = calc_sla_dt(prazo, d_ab, mun, uf)
            dentro = (d_fc <= sla_dt) if sla_dt else False
            eh_exp = 'EXPRESSO' in desc.upper()
            if dentro:
                if eh_exp:
                    if aplic == 'IPF': classif, status_prazo = 'Normal', 'FP'
                    else: classif, status_prazo = 'Expresso', 'NP'
                else:
                    classif, status_prazo = 'Normal', 'NP'
            else:
                if aplic in ('ABF', 'RFT'):
                    classif = 'Expresso' if eh_exp else 'Normal'
                    status_prazo = 'NP'
                else:
                    classif = 'Normal'
                    status_prazo = 'FP'
    elif ref_prazo and 'FISCALIZA' in ref_prazo.upper():
        if not d_fc: log = 'D. fechamento vazio; '
        elif not pz_comb: classif, status_prazo = 'Normal', 'NP'
        elif pz_comb >= d_fc: classif, status_prazo = 'Normal', 'NP'
        else: classif, status_prazo = 'Normal', 'FP'
    else:
        log = 'Ref. prazo nao reconhecida; '
        
    counts[contrato]['total'] += 1
    status_dist[contrato][status_prazo] += 1
    
    if situacao == 'CO':
        if ref_fdm == 1:
            if status_prazo in ('NP', 'FP'): counts[contrato]['TSPE1'] += 1
            if status_prazo == 'NP': counts[contrato]['TSPR1'] += 1
        elif ref_fdm == 2:
            if status_prazo in ('NP', 'FP'): counts[contrato]['TSPE2'] += 1
            if status_prazo == 'NP': counts[contrato]['TSPR2'] += 1

print('Simulated Correct TSPE and TSPR Results:')
for cont, d in counts.items():
    iapfarq1 = (d['TSPR1'] / d['TSPE1'] * 100) if d['TSPE1'] > 0 else 100.0
    iapfarq2 = (d['TSPR2'] / d['TSPE2'] * 100) if d['TSPE2'] > 0 else 100.0
    print(f'Contrato: {cont} (Total linhas: {d["total"]})')
    print(f'  Distribuicao de Status Prazo: {dict(status_dist[cont])}')
    print(f'  Servico 1 (Ref. FDM = 1): TSPE1 = {d["TSPE1"]}, TSPR1 = {d["TSPR1"]}, IAPFARQ1 = {iapfarq1:.2f}%')
    print(f'  Servico 2 (Ref. FDM = 2): TSPE2 = {d["TSPE2"]}, TSPR2 = {d["TSPR2"]}, IAPFARQ2 = {iapfarq2:.2f}%')
